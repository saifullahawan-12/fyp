from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def generate_report():
    print("Generating FYP 60% Report...")
    doc = Document()
    
    # Add Title Page
    title = doc.add_heading('EARLY DETECTION OF MALICIOUS OWNERSHIP TRANSACTIONS IN BROWSER EXTENSIONS VIA TEMPORAL BEHAVIOR DRIFT MODELING', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\n')
    subtitle = doc.add_paragraph('Final Year Project - 60% Submission Report')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # List of chapters to include
    chapters = [
        "chapter1_introduction.md",
        "chapter2_literature_review.md",
        "chapter3_methodology.md",
        "chapter4_implementation.md"
    ]
    
    reports_dir = os.path.join("..", "reports")
    
    for chapter_file in chapters:
        filepath = os.path.join(reports_dir, chapter_file)
        if not os.path.exists(filepath):
            print(f"Warning: {chapter_file} not found. Skipping.")
            continue
            
        with open(filepath, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Basic Markdown Parsing for Headings
            if line.startswith('## '):
                doc.add_heading(line[3:], level=2)
                
                # Check if this is Section 4.3, so we can insert the images right under it!
                if "4.3 Visualizing Behavioral Drift" in line:
                    doc.add_paragraph("The following figures demonstrate the output of our Isolation Forest model:")
                    
                    # Insert Anomaly Score Distribution Image
                    img1_path = os.path.join(reports_dir, "anomaly_scores_dist.png")
                    if os.path.exists(img1_path):
                        doc.add_picture(img1_path, width=Inches(5.0))
                        caption1 = doc.add_paragraph("Figure 4.1: Distribution of Anomaly Scores (Lower scores indicate higher probability of being an anomaly).")
                        caption1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                    # Insert Scatter Plot Image
                    img2_path = os.path.join(reports_dir, "anomaly_scatter.png")
                    if os.path.exists(img2_path):
                        doc.add_picture(img2_path, width=Inches(5.0))
                        caption2 = doc.add_paragraph("Figure 4.2: Scatter Plot showing the Isolation Forest successfully isolating high-permission, high-network-request extensions as suspicious (Red).")
                        caption2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('- **') or line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. '):
                doc.add_paragraph(line[3:], style='List Number')
            else:
                # Remove bold asterisks for clean text
                clean_text = line.replace('**', '')
                doc.add_paragraph(clean_text)
                
        # Add page break after each chapter
        doc.add_page_break()
        
    # Final Output
    output_path = os.path.join(reports_dir, "FYP_60_Percent_Report.docx")
    doc.save(output_path)
    print(f"Success! Professional Word Document saved to: {output_path}")

if __name__ == "__main__":
    generate_report()
